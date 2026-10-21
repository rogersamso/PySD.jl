"""Write test/data/sample.xlsx, the Excel fixture for the XLSX reader tests.

Run from the repository root:  python test/data/make_sample_xlsx.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

wb = Workbook()
ws = wb.active
ws.title = "Data"

# Scalars
ws["B2"] = 3.5
ws["C2"] = 7.0
# Row vector, column vector, 2x3 matrix
for j, v in enumerate([1.0, 2.0, 3.0]):
    ws.cell(row=4, column=2 + j, value=v)            # B4:D4
for i, v in enumerate([10.0, 20.0, 30.0]):
    ws.cell(row=2 + i, column=6, value=v)            # F2:F4
for i, row in enumerate([[1.0, 2.0, 3.0], [4.0, 5.0, 6.0]]):
    for j, v in enumerate(row):
        ws.cell(row=6 + i, column=2 + j, value=v)    # B6:D7
# Series in row mode: times in row 10, values in row 11 (from column B)
for j, (t, v) in enumerate(zip([0, 1, 2, 3, 4], [0, 10, 20, 30, 40])):
    ws.cell(row=10, column=2 + j, value=float(t))
    ws.cell(row=11, column=2 + j, value=float(v))
# Series in column mode: times in column H, values in I and J (from row 2)
for i, (t, v, w) in enumerate(zip([0, 1, 2, 3, 4], [5, 4, 3, 2, 1], [1, 2, 3, 4, 5])):
    ws.cell(row=2 + i, column=8, value=float(t))
    ws.cell(row=2 + i, column=9, value=float(v))
    ws.cell(row=2 + i, column=10, value=float(w))

names = {
    "scalar_param": "Data!$B$2",
    "row_param": "Data!$B$4:$D$4",
    "col_param": "Data!$F$2:$F$4",
    "mat_param": "Data!$B$6:$D$7",
    "xs": "Data!$H$2:$H$6",
    "ys": "Data!$I$2:$I$6",
    "ys2": "Data!$I$2:$J$6",
}
for name, ref in names.items():
    wb.defined_names[name] = DefinedName(name, attr_text=ref)
# Sheet-scoped name
ws.defined_names["local_param"] = DefinedName("local_param", attr_text="Data!$C$2")

wb.save(Path(__file__).with_name("sample.xlsx"))
