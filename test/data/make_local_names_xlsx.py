"""Write test/data/local_names.xlsx: two sheets with a named range of the
same name, local to each (as a scenario sheet and its older copy), and sheet
ids that are not the sheets' positions, as Excel leaves them after sheets are
added and moved.

Run from the repository root:  python test/data/make_local_names_xlsx.py
"""
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

path = Path(__file__).with_name("local_names.xlsx")
wb = Workbook()
readme = wb.active
readme.title = "Readme"
wb.create_sheet("Info")
new = wb.create_sheet("NZP")
old = wb.create_sheet("NZP_old")
new["C42"] = 0.086
old["C42"] = 0.13
new.defined_names["p_growth"] = DefinedName("p_growth", attr_text="NZP!$C$42")
old.defined_names["p_growth"] = DefinedName("p_growth", attr_text="NZP_old!$C$42")
wb.save(path)

# NZP is the 3rd sheet with id 5, NZP_old has id 3 (as in a real scenario file)
with zipfile.ZipFile(path) as z:
    files = {n: z.read(n) for n in z.namelist()}
xml = files["xl/workbook.xml"].decode()
for name, sid in (("NZP", 5), ("NZP_old", 3)):
    xml = re.sub(rf'(<sheet [^>]*name="{name}"[^>]*sheetId=")\d+', rf"\g<1>{sid}", xml)
files["xl/workbook.xml"] = xml.encode()
with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
    for n, data in files.items():
        z.writestr(n, data)
